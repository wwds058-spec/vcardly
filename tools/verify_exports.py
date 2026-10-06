#!/usr/bin/env python3
"""Opens the CSV, Excel and PDF files written by the iOS export tests with standard tools (csv, openpyxl) and checks
what a user would see. Usage: verify_exports.py <dir>   (needs: pip install openpyxl)"""
import csv
import io
import os
import sys

import openpyxl

d = sys.argv[1]

with open(os.path.join(d, "vcardly-contacts.csv"), "rb") as f:
    raw = f.read()
assert raw.startswith(b"\xef\xbb\xbf"), "CSV must start with a UTF-8 BOM"
rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline="")))
assert len(rows) == 4 and len(rows[0]) == 15, rows[0]
assert rows[1][0] == "Asha Rao" and rows[1][2] == "Acme, Inc" and rows[1][13] == 'Line one\nline "two"', rows[1]
assert rows[2][0].startswith("'="), "formula must be neutralised"
assert rows[3][0] == "रमेश कुमार" and rows[3][3] == "-123 456"
print("CSV OK:", len(rows) - 1, "contacts")

wb = openpyxl.load_workbook(os.path.join(d, "vcardly-workbook.xlsx"))
assert wb.sheetnames == ["Contacts", "Follow-ups"], wb.sheetnames
ws = wb["Contacts"]
assert ws["A1"].value == "Full name" and ws["A1"].font.b, "bold header"
assert ws["A2"].value == "Asha Rao" and ws["C2"].value == "Acme, Inc"
assert ws["A3"].value == '=HYPERLINK("http://x")' and ws["A3"].data_type == "s", "kept as text, never a formula"
assert ws["C3"].value == "<b>Globex</b> & Co"
assert ws.freeze_panes == "A2"
fu = wb["Follow-ups"]
assert fu["C2"].value == "Send quote" and fu["E2"].value == "Completed"
print("XLSX OK:", ws.max_row - 1, "contacts,", fu.max_row - 1, "follow-ups")

with open(os.path.join(d, "vcardly-report.pdf"), "rb") as f:
    pdf = f.read()
assert pdf.startswith(b"%PDF-") and pdf.rstrip().endswith(b"%%EOF") and b"/Type /Page" in pdf
print("PDF OK:", len(pdf), "bytes")
